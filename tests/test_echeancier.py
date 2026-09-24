from dataclasses import replace
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from echeancier import (
    ajuster_sur_solde,
    calculer,
    MODELE_PENNYLANE,
    Deblocage,
    ParametresPret,
    calculer_echeancier,
    comparer,
    exporter_pennylane,
    lire_grand_livre,
)

GRAND_LIVRE = Path(__file__).parent / "grand_livre_exemple.xlsx"


def pret_simple():
    return ParametresPret(
        capital=24_000,
        taux=4.17,
        nb_echeances=60,
        date_premier_paiement=date(2026, 1, 5),
        jour_prelevement=10,
        deblocages=[Deblocage(date(2025, 12, 10), 24_000)],
    )


def pret_banque(echeance_imposee=None, constates=()):
    deblocages, _ = lire_grand_livre(GRAND_LIVRE)
    return ParametresPret(
        capital=24_000,
        taux=4.17,
        nb_echeances=60,
        nb_echeances_differe=3,
        date_premier_paiement=date(2026, 1, 5),
        jour_prelevement=5,
        deblocages=deblocages,
        echeance_imposee=echeance_imposee,
        remboursements_constates=list(constates),
    )


def test_reproduit_le_fichier_type_pennylane():
    """Sans déblocages multiples, on retrouve à l'identique l'échéancier généré par Pennylane."""
    calcule = calculer_echeancier(pret_simple()).values.tolist()
    reference = [[c.value for c in ligne] for ligne in openpyxl.load_workbook(MODELE_PENNYLANE).active.iter_rows(min_row=6)]
    assert len(calcule) == len(reference) == 60
    for ligne, attendu in zip(calcule, reference):
        assert ligne[0] == attendu[0].date()
        assert ligne[1:] == pytest.approx(attendu[1:], abs=0.001)


def test_lecture_grand_livre():
    deblocages, remboursements = lire_grand_livre(GRAND_LIVRE)
    assert len(deblocages) == 5
    assert sum(d.montant for d in deblocages) == pytest.approx(24_000)
    assert remboursements["Capital remboursé (€)"].tolist()[0] == pytest.approx(384.40)


def test_differe_capitalise():
    r = calculer(pret_banque())
    e = r.echeancier
    # Différé : intérêts sur les fonds débloqués, ajoutés au capital, rien n'est prélevé.
    assert e["Intérêt (€)"].iloc[0] == pytest.approx(18_319 * 0.0417 * 24 / 365, abs=0.01)
    assert e["Amortissement (€)"].iloc[:3].tolist() == pytest.approx((-e["Intérêt (€)"].iloc[:3]).tolist())
    assert e["Échéance (€)"].iloc[:3].tolist() == [0, 0, 0]
    assert r.capital_amorti == pytest.approx(24_000 + r.interets_capitalises_calcules)
    # Solde au 05/03/2026 : fonds réellement versés (21 672,72 €) + intérêts ajoutés au capital.
    assert e["Solde (€)"].iloc[2] == pytest.approx(21_672.72 + r.interets_capitalises_calcules)
    assert e["Solde (€)"].iloc[-1] == pytest.approx(0)
    # Échéances constantes une fois tous les fonds versés (dernier déblocage le 24/06/2026).
    assert e["Échéance (€)"].iloc[7:-1].nunique() == 1


def test_differe_paye():
    p = pret_banque()
    p.interets_differe_capitalises = False
    e = calculer_echeancier(p)
    assert e["Amortissement (€)"].iloc[:3].tolist() == [0, 0, 0]
    assert e["Échéance (€)"].iloc[0] == e["Intérêt (€)"].iloc[0] > 0
    assert e["Amortissement (€)"].sum() == pytest.approx(24_000)


# Extrait du tableau d'amortissement bancaire du prêt n°141 : rang -> (intérêts, capital amorti, CRD).
TABLEAU_BANQUE = {
    10: (75.96, 392.49, 21_466.71),
    11: (74.60, 393.85, 21_072.86),
    16: (67.71, 400.74, 19_082.97),
    20: (62.11, 406.34, 17_466.02),
    30: (47.76, 420.69, 13_324.11),
    40: (32.91, 435.54, 9_035.98),
    50: (17.54, 450.91, 4_596.50),
    55: (9.65, 458.80, 2_318.33),
    58: (4.85, 463.60, 932.35),
    59: (3.24, 465.21, 467.14),
}


def test_reproduit_le_tableau_bancaire():
    """Échéance 468,45 € et capital de départ 24 185,76 € (21 859,20 + 2 326,56) : tableau au centime."""
    p = pret_banque(468.45)
    p.interets_capitalises_imposes = 185.76
    e = calculer_echeancier(p)
    assert e["Solde (€)"].iloc[8] == pytest.approx(21_859.20)
    assert e["Amortissement (€)"].iloc[3:9].sum() == pytest.approx(2_326.56)
    for rang, attendu in TABLEAU_BANQUE.items():
        ligne = e.iloc[rang - 1]
        assert (ligne["Intérêt (€)"], ligne["Amortissement (€)"], ligne["Solde (€)"]) == pytest.approx(attendu, abs=0.001)
    assert e["Solde (€)"].iloc[-1] == pytest.approx(0)
    _, remboursements = lire_grand_livre(GRAND_LIVRE)
    assert comparer(e, remboursements)["Écart (€)"].abs().max() == 0


def test_calage_sur_le_grand_livre():
    """Sans le tableau bancaire : capital déduit de l'échéance et des remboursements comptabilisés."""
    _, remboursements = lire_grand_livre(GRAND_LIVRE)
    r = calculer(pret_banque(468.45, remboursements["Capital remboursé (€)"]))
    assert comparer(r.echeancier, remboursements)["Écart (€)"].abs().max() == 0
    assert set(r.echeancier["Échéance (€)"].iloc[7:-1]) == {468.45}
    assert abs(r.capital_amorti - 24_185.76) < 0.2


def test_option_jours_exacts():
    p = pret_banque(468.45)
    p.interets_jours_exacts = True
    e = calculer_echeancier(p)
    avril, mai = e.iloc[3], e.iloc[4]  # 05/03 → 05/04 : 31 jours ; 05/04 → 05/05 : 30 jours
    assert avril["Intérêt (€)"] == pytest.approx(e["Solde (€)"].iloc[2] * 0.0417 * 31 / 365, abs=0.01)
    assert mai["Intérêt (€)"] == pytest.approx(avril["Solde (€)"] * 0.0417 * 30 / 365, abs=0.01)


def test_ajustement_sur_le_capital_restant_du_de_la_banque():
    """CRD de la banque au 08/09/2026 (21 859,20 €) : l'application retrouve seule le tableau bancaire."""
    a = ajuster_sur_solde(pret_banque(468.45), date(2026, 9, 8), 21_859.20)
    assert a.levier == "intérêts capitalisés"
    assert a.date_echeance == date(2026, 9, 5)
    assert a.solde_obtenu == pytest.approx(21_859.20)
    assert a.parametres.interets_capitalises_imposes == pytest.approx(185.76)
    e = calculer_echeancier(a.parametres)
    for rang, attendu in TABLEAU_BANQUE.items():
        ligne = e.iloc[rang - 1]
        assert (ligne["Intérêt (€)"], ligne["Amortissement (€)"], ligne["Solde (€)"]) == pytest.approx(attendu, abs=0.001)


def test_solde_reel_egal_au_grand_livre_plus_interets_capitalises():
    """Le solde suit les déblocages : il égale à chaque date le solde du compte 164 + 185,76 € capitalisés."""
    p = pret_banque(468.45)
    p.interets_capitalises_imposes = 185.76
    e = calculer_echeancier(p)
    grand_livre = {date(2026, 4, 5): 21_288.32, date(2026, 5, 5): 20_902.58, date(2026, 6, 5): 21_235.93,
                   date(2026, 7, 5): 22_454.35, date(2026, 8, 5): 22_064.57, date(2026, 9, 5): 21_673.44}
    for jour, solde_164 in grand_livre.items():
        assert e.loc[e["Date"] == jour, "Solde (€)"].iloc[0] == pytest.approx(solde_164 + 185.76)


def test_ajustement_sans_differe_recalcule_l_echeance():
    p = pret_simple()
    p.echeance_imposee = 400.0  # échéance erronée
    a = ajuster_sur_solde(p, date(2026, 12, 15), 19_591.09)  # CRD du fichier type Pennylane au 10/12/2026
    assert a.levier == "échéance"
    assert a.parametres.echeance_imposee == pytest.approx(443.84)
    assert a.solde_obtenu == pytest.approx(19_591.09)


def test_ajustement_avant_la_premiere_echeance():
    assert ajuster_sur_solde(pret_simple(), date(2025, 12, 1), 24_000) is None


def pret_partiel(**options):
    """Prêt de 10 000 € sur 80 mois à 4 %, dont 7 500 € seulement débloqués."""
    valeurs = dict(
        capital=10_000,
        taux=4.0,
        nb_echeances=80,
        date_premier_paiement=date(2026, 1, 10),
        jour_prelevement=10,
        montant_debloque=7_500,
    )
    return ParametresPret(**{**valeurs, **options})


def test_partiel_echeance_maintenue_solde_le_capital_plus_tot():
    r = calculer(pret_partiel(partiel_duree_reduite=True))
    e = r.echeancier
    assert r.echeance_constante == pytest.approx(142.61)  # échéance du prêt complet de 10 000 €
    assert len(e) == 80  # la durée saisie est conservée
    assert r.rang_capital_solde == 58  # capital soldé à la 58e échéance
    assert set(e["Échéance (€)"].iloc[:57]) == {142.61}
    assert (e.iloc[58:][["Intérêt (€)", "Amortissement (€)", "Échéance (€)", "Solde (€)"]] == 0).all().all()
    assert e["Amortissement (€)"].sum() == pytest.approx(7_500)


def test_partiel_l_assurance_continue_apres_le_solde_du_capital():
    e = calculer_echeancier(pret_partiel(assurance_mensuelle=3.0))
    assert len(e) == 80
    assert set(e["Assurance (€)"]) == {3.00}  # coût mensuel saisi, sur toute la durée
    assert (e["Échéance (€)"].iloc[58:] == 3.00).all()


def test_assurance_en_pourcentage_du_capital_restant_du():
    e = calculer_echeancier(pret_partiel(assurance_taux_crd=0.36))
    # 1re échéance : 7 500 € restant dû x 0,36 % / 12
    assert e["Assurance (€)"].iloc[0] == pytest.approx(2.25)
    # 2e échéance : sur le capital restant dû après la 1re
    assert e["Assurance (€)"].iloc[1] == pytest.approx(round(e["Solde (€)"].iloc[0] * 0.0036 / 12, 2))
    # L'assurance diminue avec le capital et s'arrête une fois celui-ci soldé (58e échéance).
    assert e["Assurance (€)"].is_monotonic_decreasing
    assert (e["Assurance (€)"].iloc[58:] == 0).all()


def test_assurance_en_pourcentage_trimestrielle_et_export(tmp_path):
    p = pret_partiel(assurance_taux_crd=0.36, periodicite="Trimestrielle", nb_echeances=27)
    e = calculer_echeancier(p)
    assert e["Assurance (€)"].iloc[0] == pytest.approx(6.75)  # 7 500 x 0,36 % x 3 / 12
    fichier = tmp_path / "assurance.xlsx"
    fichier.write_bytes(exporter_pennylane(p, e))
    ws = openpyxl.load_workbook(fichier).active
    assert ws["C2"].value == 0.36
    assert ws["D2"].value == pytest.approx(round(e["Assurance (€)"].sum(), 2))


def test_assurance_mensuelle_ramenee_a_la_periodicite():
    e = calculer_echeancier(pret_partiel(assurance_mensuelle=12.5, periodicite="Trimestrielle", nb_echeances=27))
    assert set(e["Assurance (€)"]) == {37.50}  # 3 mois x 12,50 €
    assert e["Échéance (€)"].iloc[0] == pytest.approx(e["Intérêt (€)"].iloc[0] + e["Amortissement (€)"].iloc[0] + 37.50)


def test_partiel_duree_maintenue_reduit_l_echeance():
    r = calculer(pret_partiel(partiel_duree_reduite=False))
    e = r.echeancier
    assert len(e) == 80
    assert r.rang_capital_solde == 80
    assert r.echeance_constante == pytest.approx(106.96)  # 7 500 € sur 80 mois
    assert e["Amortissement (€)"].sum() == pytest.approx(7_500)
    assert e["Solde (€)"].iloc[-1] == pytest.approx(0)


def test_partiel_echeance_bancaire_decide_de_la_duree():
    assert calculer(pret_partiel(echeance_imposee=142.61)).rang_capital_solde == 58
    assert calculer(pret_partiel(echeance_imposee=106.96)).rang_capital_solde == 80


def test_partiel_export_porte_sur_le_montant_debloque(tmp_path):
    p = pret_partiel()
    fichier = tmp_path / "partiel.xlsx"
    fichier.write_bytes(exporter_pennylane(p, calculer_echeancier(p)))
    ws = openpyxl.load_workbook(fichier).active
    assert ws["A2"].value == 7_500
    assert ws["G2"].value == 80
    assert ws.max_row == 5 + 80


@pytest.mark.parametrize("periodicite, mois, nombre", [("Mensuelle", 1, 80), ("Trimestrielle", 3, 27), ("Semestrielle", 6, 14), ("Annuelle", 12, 7)])
@pytest.mark.parametrize("duree_reduite", [True, False])
def test_periodicites_en_deblocage_partiel(periodicite, mois, nombre, duree_reduite):
    e = calculer_echeancier(pret_partiel(periodicite=periodicite, nb_echeances=nombre, partiel_duree_reduite=duree_reduite))
    assert (e["Date"].iloc[1] - e["Date"].iloc[0]).days in range(28 * mois, 31 * mois + 1)
    assert e["Amortissement (€)"].sum() == pytest.approx(7_500)
    assert e["Solde (€)"].iloc[-1] == pytest.approx(0)
    assert len(e) == nombre
    solde_avant_terme = (e["Solde (€)"].iloc[:-1] == 0).any()
    assert solde_avant_terme if duree_reduite else not solde_avant_terme


def pret_deux_deblocages(**options):
    """Prêt de 10 000 € sur 24 mois à 4 % : 6 000 € versés le 15/12/2025, 4 000 € le 20/03/2026."""
    valeurs = dict(
        capital=10_000,
        taux=4.0,
        nb_echeances=24,
        date_premier_paiement=date(2026, 1, 10),
        jour_prelevement=10,
        deblocages=[Deblocage(date(2025, 12, 15), 6_000), Deblocage(date(2026, 3, 20), 4_000)],
    )
    return ParametresPret(**{**valeurs, **options})


def test_echeances_totales_par_defaut():
    e = calculer_echeancier(pret_deux_deblocages())
    # Capital amorti selon le tableau du prêt complet (10 000 € sur 24 mois) dès la 1re échéance…
    theorique = calculer_echeancier(pret_deux_deblocages(deblocages=[Deblocage(date(2025, 12, 10), 10_000)]))
    assert e["Amortissement (€)"].tolist() == theorique["Amortissement (€)"].tolist()
    # … mais intérêts et solde sur les seuls fonds versés : 6 000 € jusqu'au 20/03/2026.
    assert e["Intérêt (€)"].iloc[0] == pytest.approx(6_000 * 0.04 * 26 / 365, abs=0.01)
    assert e["Solde (€)"].iloc[0] == pytest.approx(6_000 - e["Amortissement (€)"].iloc[0])
    assert e["Échéance (€)"].iloc[0] < 434.25
    assert set(e["Échéance (€)"].iloc[4:-1]) == {434.25}  # tout est versé : échéance du prêt complet


def test_echeances_proratisees_jusqu_au_dernier_deblocage():
    e = calculer_echeancier(pret_deux_deblocages(echeance_proratisee=True))
    # Tant que 6 000 € sur 10 000 € sont versés : 60 % de l'échéance.
    assert e["Échéance (€)"].iloc[:3].tolist() == [260.55] * 3
    # 1re échéance : intérêts sur les 6 000 € du 15/12/2025 au 10/01/2026 (26 jours).
    assert e["Intérêt (€)"].iloc[0] == pytest.approx(6_000 * 0.04 * 26 / 365, abs=0.01)
    # Échéance d'avril : tout est versé, échéance entière ; intérêts des 4 000 € au prorata (21 jours).
    assert e["Échéance (€)"].iloc[3] == pytest.approx(434.25)
    # Ensuite : échéance recalculée pour solder le prêt au terme.
    assert len(set(e["Échéance (€)"].iloc[4:-1])) == 1
    assert len(e) == 24
    assert e["Amortissement (€)"].sum() == pytest.approx(10_000)
    assert e["Solde (€)"].iloc[-1] == pytest.approx(0)


def test_proratisation_sans_effet_si_tout_est_verse_avant_la_1re_echeance():
    p = pret_deux_deblocages(echeance_proratisee=True)
    p.deblocages = [Deblocage(date(2025, 12, 1), 6_000), Deblocage(date(2025, 12, 20), 4_000)]
    assert calculer_echeancier(p).equals(calculer_echeancier(replace(p, echeance_proratisee=False)))


def test_proratisation_sur_le_pret_141():
    """Prêt n°141 : les déblocages de juin interviennent après la 1re échéance d'amortissement (05/04/2026)."""
    p = pret_banque(468.45)
    p.interets_capitalises_imposes = 185.76
    totale = calculer_echeancier(p)
    proratisee = calculer_echeancier(replace(p, echeance_proratisee=True))
    assert proratisee["Échéance (€)"].iloc[3] < totale["Échéance (€)"].iloc[3]  # avril : fonds pas tous versés
    assert proratisee["Amortissement (€)"].sum() == pytest.approx(totale["Amortissement (€)"].sum())
    assert proratisee["Solde (€)"].iloc[-1] == pytest.approx(0)


def test_export_respecte_le_fichier_type(tmp_path):
    p = pret_banque()
    fichier = tmp_path / "export.xlsx"
    fichier.write_bytes(exporter_pennylane(p, calculer_echeancier(p)))
    ws = openpyxl.load_workbook(fichier).active
    modele = openpyxl.load_workbook(MODELE_PENNYLANE).active
    for ligne in (1, 5):
        assert [c.value for c in ws[ligne]] == [c.value for c in modele[ligne]]
    assert [c.value for c in ws[2]] == [24_000, 4.17, 0, 0, 0, 0, 60]
    assert ws.max_row == 65
    assert ws["A6"].number_format == "dd/mm/yyyy"
