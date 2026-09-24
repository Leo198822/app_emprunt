"""Calcul d'un échéancier d'emprunt à déblocages multiples, au format Pennylane.

Principe (constaté sur le grand livre d'un prêt Crédit Agricole à déblocages successifs) :
- pendant le **différé**, les intérêts courent sur les fonds réellement débloqués : période pleine
  (taux / 12) pour les fonds déjà versés, jours exacts / 365 pour ceux versés en cours de période ; ils sont soit payés à chaque échéance, soit **capitalisés** (ajoutés au capital) ;
- à la fin du différé, la banque amortit le capital total (+ intérêts capitalisés) par
  **échéances constantes**, comme un prêt classique, quelles que soient les dates des derniers déblocages ;
- pendant l'amortissement, les intérêts sont calculés à taux / 12 sur le capital restant dû
  (vérifié au centime sur le tableau bancaire) ; une option permet les jours exacts / 365 ;
- **déblocages après la 1re échéance d'amortissement** : capital amorti selon le tableau du prêt complet
  (banque du prêt n°141), ou échéances proratisées aux fonds versés jusqu'au dernier déblocage ; dans
  les deux cas, intérêts et solde portent sur le capital réellement versé à chaque date ;
- **déblocage partiel** : l'échéancier porte sur le montant débloqué, soit avec l'échéance du prêt
  complet (le capital est soldé plus tôt ; les échéances suivantes restent jusqu'au terme, à 0 hors
  assurance), soit avec une échéance réduite étalée sur toute la durée.
"""

from __future__ import annotations

import calendar
from dataclasses import dataclass, field, replace
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd

MODELE_PENNYLANE = Path(__file__).parent / "modeles" / "Echeancier_type.xlsx"

BASE_JOURS = 365  # prorata des jours : base 365
FORMAT_DATE = "dd/mm/yyyy"  # dates au format français dans le fichier exporté
PERIODICITES = {"Mensuelle": 1, "Trimestrielle": 3, "Semestrielle": 6, "Annuelle": 12}
COLONNES = ["Date", "Intérêt (€)", "Assurance (€)", "Autres frais (€)", "Amortissement (€)", "Échéance (€)", "Solde (€)"]


@dataclass
class Deblocage:
    date: date
    montant: float


@dataclass
class ParametresPret:
    capital: float  # montant emprunté (offre de prêt)
    taux: float  # taux nominal annuel en %
    nb_echeances: int  # nombre total d'échéances, différé inclus
    date_premier_paiement: date
    jour_prelevement: int
    deblocages: list[Deblocage] = field(default_factory=list)
    nb_echeances_differe: int = 0  # échéances de différé avant l'amortissement
    interets_differe_capitalises: bool = True  # False : intérêts du différé payés à chaque échéance
    interets_capitalises_imposes: float | None = None  # montant du tableau bancaire, si connu
    periodicite: str = "Mensuelle"
    type_remboursement: str = "Échéances constantes"  # ou "Amortissement constant"
    echeance_imposee: float | None = None  # échéance hors assurance de l'offre bancaire, si connue
    remboursements_constates: list[float] = field(default_factory=list)  # capital remboursé (grand livre)
    interets_jours_exacts: bool = False  # amortissement : taux / 12 (banque) ; True : jours exacts / 365
    assurance_mensuelle: float = 0.0  # montant fixe : coût mensuel de l'assurance, jusqu'au terme du prêt
    assurance_taux_crd: float = 0.0  # ou taux annuel (%) appliqué au capital restant dû en début de période
    autres_frais: float = 0.0
    autres_frais_en_pourcentage: bool = False  # True : % du capital ; False : montant total
    montant_debloque: float | None = None  # déblocage partiel : montant réellement versé (None = en totalité)
    partiel_duree_reduite: bool = True  # partiel : True = échéance du prêt complet, durée raccourcie ;
    # False = durée maintenue, échéance recalculée sur le montant débloqué
    echeance_proratisee: bool = False  # déblocages après la 1re échéance : échéances réduites au prorata des
    # fonds versés jusqu'au dernier déblocage (False = échéances calculées sur la totalité du prêt)

    @property
    def capital_effectif(self) -> float:
        """Capital réellement débloqué, sur lequel porte l'échéancier."""
        return self.capital if self.montant_debloque is None else self.montant_debloque

    @property
    def partiel(self) -> bool:
        return self.montant_debloque is not None and self.montant_debloque < self.capital - 0.005

    @property
    def mois_par_periode(self) -> int:
        return PERIODICITES[self.periodicite]

    @property
    def taux_periodique(self) -> float:
        return self.taux / 100 * self.mois_par_periode / 12

    @property
    def nb_echeances_amortissement(self) -> int:
        return self.nb_echeances - self.nb_echeances_differe

    @property
    def assurance_par_echeance(self) -> float:
        """Coût mensuel de l'assurance ramené à la périodicité (× 3 en trimestriel, × 12 en annuel…)."""
        return round(self.assurance_mensuelle * self.mois_par_periode, 2)

    def assurance_sur(self, capital_restant_du: float) -> float:
        """Assurance d'une échéance : taux sur le capital restant dû, ou montant fixe."""
        if self.assurance_taux_crd:
            return round(max(capital_restant_du, 0.0) * self.assurance_taux_crd / 100 * self.mois_par_periode / 12, 2)
        return self.assurance_par_echeance

    @property
    def montant_total_autres_frais(self) -> float:
        if self.autres_frais_en_pourcentage:
            return round(self.capital * self.autres_frais / 100, 2)
        return round(self.autres_frais, 2)

    @property
    def taux_assurance(self) -> float:
        return self.assurance_taux_crd  # 0 pour un montant fixe : seul le montant total figure en en-tête

    @property
    def pourcentage_autres_frais(self) -> float:
        return self.autres_frais if self.autres_frais_en_pourcentage else 0.0


@dataclass
class Resultat:
    echeancier: pd.DataFrame
    capital_amorti: float  # capital + intérêts capitalisés
    interets_capitalises_calcules: float
    interets_capitalises_retenus: float
    echeance_constante: float | None
    rang_capital_solde: int  # rang de l'échéance qui solde le capital (= nombre d'échéances sans fin anticipée)


def ajouter_mois(d: date, mois: int, jour: int | None = None) -> date:
    """Décale une date de `mois` mois en gardant le jour voulu (borné à la fin du mois)."""
    total = d.month - 1 + mois
    annee, mois_ = d.year + total // 12, total % 12 + 1
    jour = jour or d.day
    return date(annee, mois_, min(jour, calendar.monthrange(annee, mois_)[1]))


def dates_echeances(p: ParametresPret) -> list[date]:
    """Première date au jour de prélèvement à partir de la date du premier paiement, puis une par période."""
    premiere = ajouter_mois(p.date_premier_paiement, 0, p.jour_prelevement)
    if premiere < p.date_premier_paiement:
        premiere = ajouter_mois(premiere, 1, p.jour_prelevement)
    return [ajouter_mois(premiere, k * p.mois_par_periode, p.jour_prelevement) for k in range(p.nb_echeances)]


def repartir(total: float, n: int) -> list[float]:
    """Répartit un montant en n parts arrondies au centime ; la dernière absorbe l'écart."""
    if n <= 0:
        return []
    part = round(total / n, 2)
    return [part] * (n - 1) + [round(total - part * (n - 1), 2)]


def interets_differe(p: ParametresPret, dates: list[date]) -> list[float]:
    """Intérêts de chaque échéance de différé, sur les fonds réellement débloqués.

    Méthode bancaire usuelle : les fonds versés avant le début de la période portent intérêt sur une
    période pleine (taux / 12) ; ceux versés pendant la période, au prorata des jours exacts / 365
    jusqu'à l'échéance. Pour la 1re échéance, les intérêts courent depuis la date de chaque déblocage.
    """
    # Sans déblocage saisi : fonds versés en totalité une période avant la première échéance.
    deblocages = sorted(p.deblocages, key=lambda d: d.date) or [
        Deblocage(ajouter_mois(dates[0], -p.mois_par_periode, p.jour_prelevement), p.capital_effectif)
    ]
    r, t = p.taux / 100, p.taux_periodique
    interets, debut = [], None
    for fin in dates[: p.nb_echeances_differe]:
        interet = 0.0
        for d in deblocages:
            if d.date >= fin:
                continue
            if debut is not None and d.date <= debut:
                interet += d.montant * t
            else:
                interet += d.montant * r * (fin - d.date).days / BASE_JOURS
        interets.append(round(interet, 2))
        debut = fin
    return interets


def repartir_ecart(montants: list[float], total: float) -> list[float]:
    """Ajuste des montants au prorata pour atteindre un total ; la dernière ligne absorbe l'arrondi."""
    somme = sum(montants)
    if not montants or somme == 0:
        return repartir(total, len(montants))
    ajustes = [round(m * total / somme, 2) for m in montants[:-1]]
    return ajustes + [round(total - sum(ajustes), 2)]


def facteurs_interets(p: ParametresPret, dates: list[date]) -> list[float]:
    """Taux d'intérêt de chaque période d'amortissement : taux périodique (ou jours exacts / 365)."""
    if not p.interets_jours_exacts:
        return [p.taux_periodique] * p.nb_echeances_amortissement
    r, k0 = p.taux / 100, p.nb_echeances_differe
    debuts = [dates[k0 - 1] if k0 else ajouter_mois(dates[0], -p.mois_par_periode, p.jour_prelevement)] + dates[k0:-1]
    return [r * (fin - debut).days / BASE_JOURS for debut, fin in zip(debuts, dates[k0:])]


def tableau_constant(base: float, echeance: float, facteurs: list[float]) -> list[tuple[float, float]]:
    """(intérêt, amortissement) de chaque échéance, arrondis au centime ligne à ligne.

    La dernière échéance solde le capital et reste égale aux autres quand l'écart n'est qu'un arrondi
    (convention Pennylane et bancaire).
    """
    lignes, restant, n = [], base, len(facteurs)
    for k, facteur in enumerate(facteurs):
        interet = round(restant * facteur, 2)
        if k == n - 1:
            a = restant
            if abs(round(echeance - a, 2) - interet) <= 0.02:
                interet = round(echeance - a, 2)
        else:
            a = min(round(echeance - interet, 2), restant)
        lignes.append((interet, round(a, 2)))
        restant = round(restant - a, 2)
    return lignes


def non_verse(p: ParametresPret, jour: date) -> float:
    """Fonds du prêt pas encore versés à une date (0 si aucun déblocage n'est renseigné)."""
    if not p.deblocages:
        return 0.0
    return max(round(p.capital_effectif - sum(d.montant for d in p.deblocages if d.date <= jour), 2), 0.0)


def interet_reel(p: ParametresPret, crd_theorique: float, debut: date, fin: date, facteur: float) -> float:
    """Intérêts d'une période sur le capital réellement versé : fonds versés avant la période sur la période
    entière, fonds versés pendant la période au prorata des jours jusqu'à l'échéance."""
    interet = (crd_theorique - non_verse(p, debut)) * facteur
    interet += sum(d.montant * p.taux / 100 * (fin - d.date).days / BASE_JOURS for d in p.deblocages if debut < d.date <= fin)
    return round(interet, 2)


def tableau_proratise(
    p: ParametresPret, dates: list[date], base: float, echeance: float, facteurs: list[float]
) -> list[tuple[float, float]]:
    """(intérêt, amortissement) quand des fonds sont versés après le début de l'amortissement.

    Tant que tout n'est pas débloqué, l'échéance est réduite au prorata des fonds versés et les intérêts
    portent sur le capital réellement versé (prorata des jours pour un déblocage en cours de période).
    Après le dernier déblocage, l'échéance est recalculée sur le capital restant dû pour finir au terme.
    """
    k0, n = p.nb_echeances_differe, len(facteurs)
    debut = dates[k0 - 1] if k0 else ajouter_mois(dates[0], -p.mois_par_periode, p.jour_prelevement)
    lignes, restant, echeance_finale = [], base, None
    for k, (fin, facteur) in enumerate(zip(dates[k0 : k0 + n], facteurs)):
        reste_a_verser = non_verse(p, fin)
        if k == n - 1:
            interet, a = interet_reel(p, restant, debut, fin, facteur), restant
        elif reste_a_verser > 0.005 or non_verse(p, debut) > 0.005:
            # Période touchée par des fonds non encore versés : échéance et intérêts au prorata.
            interet = interet_reel(p, restant, debut, fin, facteur)
            a = round(round(echeance * (base - reste_a_verser) / base, 2) - interet, 2)
        else:
            if echeance_finale is None:
                echeance_finale = echeance_pour(restant, p.taux_periodique, n - k)
            interet = round(restant * facteur, 2)
            a = min(round(echeance_finale - interet, 2), restant)
        lignes.append((interet, round(a, 2)))
        restant = round(restant - a, 2)
        debut = fin
    return lignes


def echeance_pour(base: float, t: float, n: int) -> float:
    return round(base / n if t == 0 else base * t / (1 - (1 + t) ** -n), 2)


def calibrer_base(echeance: float, t: float, facteurs: list[float], constates: list[float]) -> float:
    """Capital amorti correspondant à une échéance connue.

    Sans données comptables : valeur actuelle des échéances. Avec les remboursements de capital du
    grand livre : milieu de la plage de capitaux (au centime) qui les reproduit exactement.
    """
    n = len(facteurs)
    base = round(echeance * n if t == 0 else echeance * (1 - (1 + t) ** -n) / t, 2)
    if not constates:
        return base
    compatibles = [
        c / 100
        for c in range(round(base * 100) - 1000, round(base * 100) + 1000)
        if [a for _, a in tableau_constant(c / 100, echeance, facteurs)[: len(constates)]] == constates
    ]
    return compatibles[len(compatibles) // 2] if compatibles else base


def calculer(p: ParametresPret) -> Resultat:
    """Échéancier complet : différé puis amortissement du capital (+ intérêts capitalisés)."""
    n, t = p.nb_echeances_amortissement, p.taux_periodique
    if n <= 0:
        raise ValueError("Le nombre d'échéances doit être supérieur au nombre d'échéances de différé.")
    dates = dates_echeances(p)
    interets = interets_differe(p, dates)
    facteurs = facteurs_interets(p, dates)
    calcules = round(sum(interets), 2) if p.interets_differe_capitalises else 0.0

    capitalises = calcules
    if p.interets_differe_capitalises and p.interets_capitalises_imposes is not None:
        capitalises = round(p.interets_capitalises_imposes, 2)
    base = round(p.capital_effectif + capitalises, 2)

    echeance = None
    proratise = (
        p.echeance_proratisee
        and p.type_remboursement == "Échéances constantes"
        and bool(p.deblocages)
        and max(d.date for d in p.deblocages) > dates[p.nb_echeances_differe]
    )
    if p.type_remboursement == "Amortissement constant":
        amortissements = repartir(base, n)
        lignes_amort, restant = [], base
        for a, facteur in zip(amortissements, facteurs):
            lignes_amort.append((round(restant * facteur, 2), a))
            restant = round(restant - a, 2)
    else:
        if p.echeance_imposee:
            echeance = round(p.echeance_imposee, 2)
            if p.interets_differe_capitalises and p.interets_capitalises_imposes is None:
                # L'échéance de la banque fixe le capital amorti, donc les intérêts capitalisés.
                if not p.partiel:
                    base = calibrer_base(echeance, t, facteurs, p.remboursements_constates)
                    capitalises = round(base - p.capital_effectif, 2)
        elif p.partiel and p.partiel_duree_reduite:
            # Échéance du prêt complet appliquée au montant débloqué : le remboursement s'arrête plus tôt.
            echeance = echeance_pour(round(p.capital + capitalises, 2), t, n)
        else:
            echeance = echeance_pour(base, t, n)
        if proratise:
            lignes_amort = tableau_proratise(p, dates, base, echeance, facteurs)
        else:
            lignes_amort = tableau_constant(base, echeance, facteurs)

    if p.interets_differe_capitalises and interets and capitalises != calcules:
        # Total imposé par la banque : l'écart est réparti au prorata des intérêts de chaque mois.
        interets = repartir_ecart(interets, capitalises)

    # Capital soldé avant la fin (déblocage partiel, échéance maintenue) : les échéances suivantes
    # restent dans l'échéancier, à 0 hors assurance et frais, jusqu'au terme prévu du prêt.
    frais = repartir(p.montant_total_autres_frais, p.nb_echeances)
    # Le tableau théorique porte sur la totalité du prêt ; le solde affiché et les intérêts portent sur le
    # capital réellement versé à chaque date (les fonds non encore débloqués n'en font pas partie).
    lignes, solde = [], p.capital_effectif
    debut = ajouter_mois(dates[0], -p.mois_par_periode, p.jour_prelevement)
    for k, d in enumerate(dates):
        crd_debut = round(solde - non_verse(p, debut), 2)
        if k < p.nb_echeances_differe:
            interet = interets[k]
            amort = -interet if p.interets_differe_capitalises else 0.0
        else:
            interet, amort = lignes_amort[k - p.nb_echeances_differe]
            fonds_en_attente = non_verse(p, debut) > 0.005 or any(debut < x.date <= d for x in p.deblocages)
            if fonds_en_attente and not proratise:
                interet = interet_reel(p, solde, debut, d, facteurs[k - p.nb_echeances_differe])
        assurance = p.assurance_sur(crd_debut)  # capital restant dû réel en début de période
        solde = round(solde - amort, 2)
        paye = round(interet + amort, 2)  # nul pendant un différé capitalisé
        solde_reel = round(solde - non_verse(p, d), 2)
        lignes.append([d, interet, assurance, frais[k], amort, round(paye + assurance + frais[k], 2), solde_reel])
        debut = d
    echeancier = pd.DataFrame(lignes, columns=COLONNES)
    soldees = [k + 1 for k in range(p.nb_echeances_differe, len(lignes)) if lignes[k][6] <= 0.005]
    return Resultat(echeancier, base, calcules, capitalises, echeance, soldees[0] if soldees else len(lignes))


def calculer_echeancier(p: ParametresPret) -> pd.DataFrame:
    return calculer(p).echeancier


@dataclass
class Ajustement:
    parametres: ParametresPret  # paramètres recalculés
    date_echeance: date  # échéance retenue (dernière à la date saisie ou avant)
    solde_obtenu: float
    levier: str  # « intérêts capitalisés » ou « échéance »


def solde_a_la_date(p: ParametresPret, jour: date) -> tuple[date, float] | None:
    """Capital restant dû après la dernière échéance à la date donnée ou avant."""
    e = calculer_echeancier(p)
    passees = e[e["Date"] <= jour]
    if passees.empty:
        return None
    return passees["Date"].iloc[-1], float(passees["Solde (€)"].iloc[-1])


def ajuster_sur_solde(p: ParametresPret, jour: date, solde_cible: float) -> Ajustement | None:
    """Recalcule l'échéancier pour que le capital restant dû à une date corresponde à celui de la banque.

    Avec un différé capitalisé, on ajuste le montant des intérêts capitalisés (donc le capital à
    amortir) ; sinon on ajuste l'échéance constante. Recherche par dichotomie au centime.
    """
    if dates_echeances(p)[0] > jour:
        return None
    if p.nb_echeances_differe and p.interets_differe_capitalises:
        levier, croissant = "intérêts capitalisés", True
        variante = lambda cts: replace(p, interets_capitalises_imposes=cts / 100)  # noqa: E731
        bas, haut = 0, round(p.capital_effectif * 100)
    elif p.type_remboursement == "Échéances constantes":
        levier, croissant = "échéance", False  # une échéance plus forte réduit le capital restant dû
        variante = lambda cts: replace(p, echeance_imposee=cts / 100)  # noqa: E731
        bas, haut = 1, round(p.capital * 100)
    else:
        return None

    def ecart(cts: int) -> float:
        return solde_a_la_date(variante(cts), jour)[1] - solde_cible

    # Plus petite valeur (en centimes) dont le solde dépasse (ou, en décroissant, passe sous) la cible.
    while bas < haut:
        milieu = (bas + haut) // 2
        if (ecart(milieu) >= 0) == croissant:
            haut = milieu
        else:
            bas = milieu + 1
    meilleur = min((c for c in (bas - 1, bas, bas + 1) if c > 0), key=lambda c: abs(ecart(c)))
    parametres = variante(meilleur)
    date_echeance, solde = solde_a_la_date(parametres, jour)
    return Ajustement(parametres, date_echeance, solde, levier)


def exporter_pennylane(p: ParametresPret, echeancier: pd.DataFrame) -> bytes:
    """Remplit le fichier type Pennylane (en-tête lignes 1-2, échéances à partir de la ligne 6)."""
    wb = openpyxl.load_workbook(MODELE_PENNYLANE)
    ws = wb.active
    format_nombre = ws["A2"].number_format
    style_ligne = [ws.cell(row=6, column=c)._style for c in range(1, 8)]
    ws.delete_rows(6, ws.max_row)

    entete = [
        p.capital_effectif,
        p.taux,
        p.taux_assurance,
        round(float(echeancier["Assurance (€)"].sum()), 2),
        p.pourcentage_autres_frais,
        p.montant_total_autres_frais,
        len(echeancier),
    ]
    for c, valeur in enumerate(entete, start=1):
        ws.cell(row=2, column=c, value=valeur)

    for i, ligne in enumerate(echeancier.itertuples(index=False), start=6):
        for c, valeur in enumerate(ligne, start=1):
            cellule = ws.cell(row=i, column=c, value=valeur)
            cellule._style = style_ligne[c - 1]
            cellule.number_format = FORMAT_DATE if c == 1 else format_nombre

    sortie = BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


def lire_grand_livre(fichier) -> tuple[list[Deblocage], pd.DataFrame]:
    """Lit un export « Grand livre » Pennylane du compte d'emprunt (colonnes Date, Débit, Crédit).

    Retourne les déblocages (crédits) et les remboursements de capital constatés (débits).
    """
    gl = pd.read_excel(fichier)
    gl["Date"] = pd.to_datetime(gl["Date"], dayfirst=True).dt.date
    gl[["Débit", "Crédit"]] = gl[["Débit", "Crédit"]].apply(pd.to_numeric, errors="coerce").fillna(0.0)
    deblocages = [Deblocage(r.Date, round(float(r.Crédit), 2)) for r in gl.itertuples() if r.Crédit > 0]
    remboursements = gl.loc[gl["Débit"] > 0, ["Date", "Débit"]].rename(columns={"Débit": "Capital remboursé (€)"})
    return deblocages, remboursements.reset_index(drop=True)


def comparer(echeancier: pd.DataFrame, remboursements: pd.DataFrame) -> pd.DataFrame:
    """Rapproche le capital remboursé en comptabilité de l'amortissement calculé, mois par mois."""
    calc = echeancier[echeancier["Amortissement (€)"] > 0][["Date", "Amortissement (€)"]].copy()
    calc["Mois"] = [d.strftime("%Y-%m") for d in calc["Date"]]
    reel = remboursements.copy()
    reel["Mois"] = [d.strftime("%Y-%m") for d in reel["Date"]]
    comp = reel.drop(columns="Date").merge(calc, on="Mois", how="left")
    comp["Écart (€)"] = (comp["Capital remboursé (€)"] - comp["Amortissement (€)"]).round(2)
    return comp[["Date", "Capital remboursé (€)", "Amortissement (€)", "Écart (€)"]]
